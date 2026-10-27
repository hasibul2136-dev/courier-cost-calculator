"""
Rate Card & Data Manager
Engineered for AMAN GRAPHICS AND DESIGNS LTD.
Handles reading, writing, updating courier rate matrices,
destination zones, exporting Work Order quotes to Excel, and syncing with original Excel file.
"""

import os
import json
import shutil
import datetime
from typing import Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
RATES_FILE = os.path.join(DATA_DIR, "rates.json")
TMP_RATES_FILE = os.path.join("/tmp", "agdl_rates.json") if os.name != 'nt' else os.path.join(os.environ.get('TEMP', DATA_DIR), "agdl_rates.json")
ORIGINAL_DESKTOP_EXCEL = r"D:\Office\OneDrive - Aman Knittings Ltd\Desktop\All Courrier Cost Calculation.xlsx"
COMPANY_NAME = "AMAN GRAPHICS AND DESIGNS LTD."


def load_rates() -> Dict[str, Any]:
    """Loads courier rate cards and app settings from JSON."""
    file_to_read = TMP_RATES_FILE if os.path.exists(TMP_RATES_FILE) else RATES_FILE
    if not os.path.exists(file_to_read):
        file_to_read = RATES_FILE

    if not os.path.exists(file_to_read):
        raise FileNotFoundError(f"Rates file not found: {file_to_read}")

    with open(file_to_read, "r", encoding="utf-8") as f:
        data = json.load(f)

    # Ensure backwards compatibility with couriers key
    if "couriers" not in data and "destination_zones" in data:
        default_zone = data.get("settings", {}).get("default_zone", "china_hk")
        data["couriers"] = data["destination_zones"].get(default_zone, {}).get("couriers", {})

    return data


def save_rates(data: Dict[str, Any]) -> bool:
    """Saves updated courier rates and settings to JSON. Supports read-only serverless."""
    try:
        os.makedirs(DATA_DIR, exist_ok=True)
        with open(RATES_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        return True
    except OSError:
        # Fallback for read-only serverless environments (Vercel)
        try:
            with open(TMP_RATES_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            return True
        except Exception:
            return False


def get_zone_data(zone_id: str) -> Dict[str, Any]:
    """Retrieves courier data for a specific destination zone."""
    data = load_rates()
    zones = data.get("destination_zones", {})
    if zone_id in zones:
        return zones[zone_id]
    # Default to china_hk
    return zones.get("china_hk", list(zones.values())[0] if zones else {})


def update_zone_courier_rate(zone_id: str, courier_code: str, courier_info: Dict[str, Any]) -> bool:
    """Updates a courier's rate card in a specific destination zone."""
    data = load_rates()
    zones = data.setdefault("destination_zones", {})
    if zone_id in zones:
        zones[zone_id].setdefault("couriers", {})[courier_code] = courier_info
        # Also update root couriers if active zone
        if data.get("settings", {}).get("default_zone") == zone_id:
            data.setdefault("couriers", {})[courier_code] = courier_info
        return save_rates(data)
    return False


def export_quote_to_excel(quote_data: Dict[str, Any], filepath: str) -> str:
    """
    Exports a professional Work Order (WO) courier cost sheet to Excel.
    Tailored for AMAN GRAPHICS AND DESIGNS LTD.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Courier Cost WO"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Dark Navy
    header_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")

    subhead_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    subhead_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    highlight_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Soft yellow
    accent_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")  # Soft green
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title block
    ws.merge_cells("A1:F1")
    ws["A1"] = "AMAN GRAPHICS AND DESIGNS LTD. - COURIER COST & WORK ORDER (WO) CALCULATION"
    ws["A1"].font = header_font
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Route & Buyer block
    origin = quote_data.get("origin", "Dhaka, Bangladesh")
    destination = quote_data.get("destination", "China")
    zone_name = quote_data.get("zone_name", "Far East")
    buyer_name = quote_data.get("buyer_name")

    ws["A3"] = "Origin (From):"
    ws["B3"] = origin
    ws["D3"] = "Destination (To):"
    ws["E3"] = f"{buyer_name} — {destination}" if buyer_name else destination

    ws["A4"] = "Generated Date:"
    ws["B4"] = datetime.datetime.now().strftime("%d-%b-%Y %I:%M %p")
    ws["D4"] = "Selected Courier:"
    ws["E4"] = quote_data.get("courier_name", quote_data.get("courier_code", "N/A"))

    ws["A5"] = "Actual Weight:"
    ws["B5"] = f"{quote_data.get('actual_weight_kg', 0)} Kg"
    ws["D5"] = "Chargeable Weight:"
    ws["E5"] = f"{quote_data.get('chargeable_weight_kg', 0)} Kg"

    ws["A6"] = "Exchange Rate:"
    ws["B6"] = f"1 USD = {quote_data.get('exchange_rate', 123.83)} BDT"
    ws["D6"] = "Tariff Zone / Slab:"
    ws["E6"] = f"{zone_name} | {quote_data.get('tier_name', 'N/A')}"

    for r in range(3, 7):
        ws[f"A{r}"].font = bold_font
        ws[f"D{r}"].font = bold_font
        ws[f"B{r}"].font = regular_font
        ws[f"E{r}"].font = regular_font

    # Breakdown Table Headers
    headers = ["Item / Component", "Description / Basis", "Rate (USD)", "Amount (USD)", "Rate in BDT (৳)", "Notes"]
    ws.append([])  # row 7
    ws.append(headers)  # row 8

    for c_idx in range(1, 7):
        cell = ws.cell(8, c_idx)
        cell.fill = subhead_fill
        cell.font = subhead_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[8].height = 25

    # Breakdown Items
    rows = [
        ("Base Freight", f"Applied slab: {quote_data.get('tier_name')}", quote_data.get("rate_applied", 0), quote_data.get("base_freight_usd", 0)),
        ("Fuel Surcharge", f"{quote_data.get('fuel_surcharge_percent', 0)}% of Base Freight", f"{quote_data.get('fuel_surcharge_percent', 0)}%", quote_data.get("fuel_surcharge_usd", 0)),
        ("Handling / Doc Fee", "Airway Bill & Documentation", quote_data.get("handling_fee_usd", 10), quote_data.get("handling_fee_usd", 10)),
        ("TOTAL REQUIRED VALUE", "Full Cost of Courier", "-", quote_data.get("total_required_usd", 0)),
        ("PER KG VALUE FOR WO", "Work Order Rate (Total / Weight)", "-", quote_data.get("per_kg_usd_for_wo", 0)),
    ]

    ex_rate = quote_data.get("exchange_rate", 123.83)

    for item, desc, rate_str, amt_usd in rows:
        amt_bdt = round(float(amt_usd) * ex_rate, 2) if isinstance(amt_usd, (int, float)) else "-"
        ws.append([item, desc, rate_str, amt_usd, amt_bdt, ""])
        current_r = ws.max_row

        is_highlight = "TOTAL" in item or "PER KG" in item
        for c in range(1, 7):
            cell = ws.cell(current_r, c)
            cell.border = thin_border
            if is_highlight:
                cell.font = bold_font
                cell.fill = highlight_fill if "TOTAL" in item else accent_fill
            else:
                cell.font = regular_font

            if c in (3, 4, 5):
                cell.alignment = Alignment(horizontal="right")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    wb.save(filepath)
    return filepath


def sync_to_desktop_excel(data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Safely synchronizes current rates back to the original Desktop Excel file.
    """
    if not os.path.exists(ORIGINAL_DESKTOP_EXCEL):
        return {"success": False, "message": f"Original file not found at {ORIGINAL_DESKTOP_EXCEL}"}

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = ORIGINAL_DESKTOP_EXCEL.replace(".xlsx", f"_backup_{timestamp}.xlsx")
    try:
        shutil.copy2(ORIGINAL_DESKTOP_EXCEL, backup_path)
    except Exception as e:
        return {"success": False, "message": f"Could not create backup: {e}"}

    try:
        wb = openpyxl.load_workbook(ORIGINAL_DESKTOP_EXCEL)
        sheet = wb["Mishu"] if "Mishu" in wb.sheetnames else wb.active

        # Use china_hk couriers for original spreadsheet
        china_zone = data.get("destination_zones", {}).get("china_hk", {})
        couriers = china_zone.get("couriers", data.get("couriers", {}))

        col_mapping = {}
        for c in range(2, sheet.max_column + 1):
            val = sheet.cell(1, c).value
            if val:
                col_mapping[str(val).strip()] = c

        for c_code, c_data in couriers.items():
            if c_code in col_mapping:
                c_idx = col_mapping[c_code]
                rates = c_data.get("rates", {})
                sheet.cell(2, c_idx).value = rates.get("first_half")
                sheet.cell(3, c_idx).value = rates.get("second_half")
                sheet.cell(4, c_idx).value = rates.get("tier_11_20")
                sheet.cell(5, c_idx).value = rates.get("tier_21_50")
                sheet.cell(6, c_idx).value = rates.get("tier_51_100")
                sheet.cell(7, c_idx).value = rates.get("tier_above_100")
                fuel_rate = c_data.get("fuel_surcharge_percent", 0.0) / 100.0
                sheet.cell(8, c_idx).value = fuel_rate

        wb.save(ORIGINAL_DESKTOP_EXCEL)
        return {
            "success": True,
            "message": "Desktop Excel file updated successfully!",
            "backup_created": backup_path
        }
    except Exception as e:
        return {"success": False, "message": f"Error updating Excel: {e}"}
